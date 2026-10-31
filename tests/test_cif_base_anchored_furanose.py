"""Explicit base anchoring replaces host sugar geometry, never target chemistry."""
import json
from unittest.mock import patch
import pytest
from openpyxl import load_workbook
from tests.test_cif_whole_furanose import setup, run as old_run
from tests.test_cif_catalogue import snapshot


def inputs(tmp_path, *, bridged=False, rename=False, hand='D'):
    paths, context, roles = setup(tmp_path, bridged=bridged, rename=rename, hand=hand)
    from nasolve.cif_stereochemistry_audit import _values
    ids = _values(paths[0], 'Z99')['_chem_comp_atom.atom_id']
    # The workbook is authoritative even when every source atom has been renamed.
    if rename:
        original = ["N1", "C2", "C6"] if bridged else ["BN", "BC", "BX"]
        plain = tmp_path/'plain'; plain.mkdir()
        other, _, _ = setup(plain, bridged=bridged, hand=hand)
        old = _values(other[0], 'Z99')['_chem_comp_atom.atom_id']
        names = [ids[old.index(n)] for n in original]
    else:names = ['N1', 'C2', 'C6'] if bridged else ['BN', 'BC', 'BX']
    book = load_workbook(paths[3]); sheet = next(s for s in book.worksheets if s.max_row > 1)
    for role, name in zip(['C2', 'C6'], names[1:]):
        column = sheet.max_column+1; sheet.cell(1,column,role); sheet.cell(2,column,name)
    book.save(paths[3]); book.close()
    lines = paths[4].read_text().splitlines(keepends=True)
    for i,line in enumerate(lines):
        if line[:6]=='ATOM  ' and line[21:26]=='A  11':
            lines[i]=line[:46]+f'{-float(line[46:54]):8.3f}'+line[54:]
    paths[4].write_text(''.join(lines))
    return paths, context


def run(paths, context, roles=None):
    from nasolve.cif_whole_furanose import prepare_base_anchored_furanose_control
    return prepare_base_anchored_furanose_control(*paths,code='Z99',context=context,
        base_frame_roles=['N1','C2','C6'] if roles is None else roles)


@pytest.mark.parametrize('bridged,rename,hand',[(False,False,'D'),(False,True,'L'),(True,False,'D'),(True,True,'D')])
def test_different_host_hands_allow_explicit_base_placement_only(tmp_path,bridged,rename,hand):
    paths,c=inputs(tmp_path,bridged=bridged,rename=rename,hand=hand)
    with pytest.raises(ValueError,match='Host and target sugar hands'):old_run(paths,c)
    before=snapshot(tmp_path)
    q=run(paths,c)
    assert q['alignment']=='NAR_BASE_FRAME' and not q['automatic_adoption']
    assert not all(r['compatible'] for r in q['host_hand_checks'])
    assert len({r['center'] for r in q['retained_hands']})==4
    assert q['max_internal_distance_error_angstrom']<=.002
    assert q['native_linkage']=='NOT_ASSESSED' and q['model_optimization']=='NOT_RUN'
    unchanged=lambda p:[l for l in p.read_text().splitlines() if not(l[:6] in {'ATOM  ','HETATM'} and l[21:26]=='A  11')]
    assert unchanged(paths[4])==unchanged(paths[-1]/'model.pdb')
    assert all(snapshot(tmp_path)[k]==v for k,v in before.items())
    assert str(tmp_path) not in json.dumps(q)


@pytest.mark.parametrize('roles',[[],['N1','C2'],['N1','C2','C2'],['C2','N1','C6'],['N1','C2','missing'],'N1,C2,C6',[1,'C2','C6']])
def test_frame_requires_explicit_distinct_glycosidic_base_roles(tmp_path,roles):
    paths,c=inputs(tmp_path)
    with pytest.raises(ValueError):run(paths,c,roles)
    assert not paths[-1].exists()


@pytest.mark.parametrize('defect',['sugar_anchor','hydrogen_anchor','mirror_candidate','flat_host','missing_host_anchor'])
def test_frame_and_target_guards_remain_fail_closed(tmp_path,defect):
    from nasolve.cif_stereochemistry_audit import _values
    from tests.test_cif_stereochemistry_audit import write
    paths,c=inputs(tmp_path)
    if defect in {'sugar_anchor','hydrogen_anchor'}:
        b=load_workbook(paths[3]);s=next(s for s in b.worksheets if s.max_row > 1)
        for cell in s[1]:
            if cell.value=='C6':s.cell(2,cell.column,'C2' if defect=='sugar_anchor' else 'C1H0')
        b.save(paths[3]);b.close()
    elif defect=='mirror_candidate':
        v=_values(paths[1],'Z99');v['_chem_comp_atom.z']=[str(-float(x)) for x in v['_chem_comp_atom.z']];write(paths[1],v)
    else:
        lines=paths[4].read_text().splitlines(keepends=True);out=[]
        for line in lines:
            if line[:6]=='ATOM  ' and line[21:26]=='A  11':
                if defect=='missing_host_anchor' and line[12:16].strip()=='BX':continue
                if defect=='flat_host':line=line[:30]+f'{0:8.3f}'*3+line[54:]
            out.append(line)
        paths[4].write_text(''.join(out))
    with pytest.raises(ValueError):run(paths,c)
    assert not paths[-1].exists()


def test_cold_replay_immutable_output_and_rehashed_role_forgery(tmp_path):
    from nasolve.cif_whole_furanose import read_base_anchored_furanose_control,read_whole_furanose_control
    from nasolve.cif_fana_component import _digest
    original=tmp_path/'original';original.mkdir();paths,c=inputs(original,bridged=True);q=run(paths,c)
    with pytest.raises(FileExistsError):run(paths,c)
    with pytest.raises(ValueError):read_whole_furanose_control(paths[-1])
    moved=original.rename(tmp_path/'moved');control=moved/'control';before=snapshot(moved)
    with patch('subprocess.Popen',side_effect=AssertionError('offline')):
        assert read_base_anchored_furanose_control(control)==q
    assert not original.exists() and snapshot(moved)==before
    selection=control/'selection.json';s=json.loads(selection.read_text());s['base_frame_roles']=['N1','C6','C2'];selection.write_text(json.dumps(s))
    receipt=control/'receipt.json';r=json.loads(receipt.read_text());r['files']['selection.json']=_digest(selection);receipt.write_text(json.dumps(r))
    with pytest.raises(ValueError):read_base_anchored_furanose_control(control)


def test_host_frame_uses_independent_workbook_atom_names(tmp_path):
    from nasolve.cif_stereochemistry_audit import _values
    from tests.test_cif_stereochemistry_audit import write
    paths,c=inputs(tmp_path,bridged=True)
    v=_values(paths[2],'Z99');mapping={n:'Y'+str(i) for i,n in enumerate(v['_chem_comp_atom.atom_id'])}
    for key in v:
        if 'atom_id' in key:v[key]=[mapping.get(n,n) for n in v[key]]
        if key.endswith('.comp_id') or key=='_chem_comp.id':v[key]=['HST']*len(v[key])
    paths[2]=write(tmp_path/'host.cif',v)
    book=load_workbook(paths[3]);sheet=next(s for s in book.worksheets if s.max_row>1)
    values=[cell.value for cell in sheet[2]]
    values[0]='HST'
    for i,cell in enumerate(sheet[1]):
        if cell.value in {'N1','C2','C5','C6'}:values[i]=mapping[values[i]]
    sheet.append(values);book.save(paths[3]);book.close()
    lines=paths[4].read_text().splitlines(keepends=True)
    for i,line in enumerate(lines):
        if line[:6]=='ATOM  ' and line[21:26]=='A  11':
            lines[i]=line[:12]+f"{mapping[line[12:16].strip()]:>4}"+line[16:17]+'HST'+line[20:]
    paths[4].write_text(''.join(lines))
    q=run(paths,c)
    assert q['target_frame_atoms']==['N1','C2','C6']
    assert q['host_frame_atoms']==[mapping[n] for n in q['target_frame_atoms']]
    assert len(q['retained_hands'])==7
