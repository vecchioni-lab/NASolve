"""Catalogue completeness, honest chemistry boundaries and read-only provenance."""
from copy import deepcopy
import json
from pathlib import Path
import shutil
from unittest.mock import patch

import openpyxl
import pytest
from restraints.residue_library import AUTHORITATIVE_SHEETS

from nasolve.cli import main

HEADERS = ['Ligand code', 'Name', 'Base Analog', 'Sugar Type', 'Phosphate',
           'Function', 'N1', 'C5', 'C5-mod']
PUBLIC = Path(__file__).parent / 'fixtures' / 'public_ccd'


def workbook(root, entries):
    path = root / 'Ligands.xlsx'
    book = openpyxl.Workbook()
    book.remove(book.active)
    for sheet in AUTHORITATIVE_SHEETS:
        page = book.create_sheet(sheet)
        page.append(HEADERS)
        for name, row in entries:
            if name == sheet:
                page.append(row)
    book.save(path)
    book.close()
    return path


def row(code='PDU', **changes):
    values = dict(zip(HEADERS, [code, 'test name', 'T', 'DNA', None, 'role', 'N1', 'C5', None]))
    values.update(changes)
    return [values[h] for h in HEADERS]


def source(root, code='PDU'):
    path = root / code[0].lower() / f'data_{code}.cif'
    path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(PUBLIC / f'{code}.cif', path)
    return path


def snapshot(root):
    return {p.relative_to(root).as_posix(): (p.read_bytes(), p.stat().st_mtime_ns)
            for p in root.rglob('*') if p.is_file()}


def test_catalogue_cli_exists_before_implementation(tmp_path, capsys):
    book = workbook(tmp_path, [('thymine', row())])
    ccd = tmp_path / 'ccd'
    source(ccd)
    with patch('nasolve.cli.discover_phenix') as discover, patch('subprocess.Popen') as process:
        assert main(['cif-doctor', 'catalogue', '--workbook', str(book),
                     '--ccd-root', str(ccd), '--json']) == 2
    report = json.loads(capsys.readouterr().out)
    assert report['summary']['distinct_codes'] == 1
    assert report['native_execution'] == 'NOT_RUN'
    discover.assert_not_called()
    process.assert_not_called()


def test_rows_aliases_blanks_and_nonstandard_backbones_are_not_lost(tmp_path):
    from nasolve.cif_catalogue import audit_catalogue
    entries = [('thymine', row()), ('thymine', row(Name='descriptive alias', **{'C5-mod': 'propynyl'})),
               ('thymine', row('ATL', **{'Sugar Type': 'LNA'})),
               ('adenine', row('APN', **{'Sugar Type': 'PNA'})),
               ('unique', row('X99', **{'Sugar Type': None}))]
    book = workbook(tmp_path, entries)
    ccd = tmp_path / 'ccd'
    source(ccd); source(ccd, 'ATL')
    before = snapshot(tmp_path)
    report = audit_catalogue(book, ccd)
    assert snapshot(tmp_path) == before
    assert report['summary']['records'] == 5
    assert report['summary']['distinct_codes'] == 4
    assert report['summary']['phosphate_labels'] == [{'value': None, 'records': 5}]
    cases = {c['component']: c for c in report['components']}
    assert cases['PDU']['record_resolution'] == 'EQUIVALENT_ROLE_RECORDS'
    assert [r['row'] for r in cases['PDU']['workbook_rows']] == [2, 3]
    assert cases['ATL']['source']['declarations']['phosphorus_atoms']
    assert cases['ATL']['source']['declarations']['op3'] is None
    assert 'OXT' in cases['ATL']['source']['declarations']['phosphorus_atoms'][0]['neighbors']
    assert cases['PDU']['source']['declarations']['op3']['element'] == 'O'
    assert cases['APN']['source']['status'] == 'MISSING'
    assert cases['APN']['workbook_rows'][0]['labels']['Sugar Type'] == 'PNA'
    assert all(c['linkage_profile'] == 'NOT_ASSIGNED' and not c['automatic_adoption'] for c in cases.values())
    assert report['summary']['native_qualification'] == {'NOT_ASSESSED': 4}
    assert not report['transformations']
    assert str(tmp_path) not in json.dumps(report)


@pytest.mark.parametrize('field,value', [('N1', 'N9'), ('Sugar Type', 'RNA'),
                                         ('Phosphate', 'no phosphate'), ('Function', 'other')])
def test_conflicting_aliases_are_never_selected(tmp_path, field, value):
    from nasolve.cif_catalogue import audit_catalogue
    book = workbook(tmp_path, [('thymine', row()), ('thymine', row(**{field: value}))])
    ccd = tmp_path / 'ccd'; source(ccd)
    case = audit_catalogue(book, ccd)['components'][0]
    assert case['record_resolution'] == 'CONFLICTING_OPERATIONAL_RECORDS'
    assert case['source']['graph_status'] == 'PASS'
    assert case['mapped_roles']['status'] == 'NOT_ASSESSED'
    assert 'RECONCILE_OPERATIONAL_RECORDS' in case['next_actions']


def test_role_name_coverage_is_not_a_recipe_or_attachment_proof(tmp_path):
    from nasolve.cif_catalogue import audit_catalogue
    book = workbook(tmp_path, [('thymine', row(N1='DOES_NOT_EXIST'))])
    ccd = tmp_path / 'ccd'; source(ccd)
    case = audit_catalogue(book, ccd)['components'][0]
    assert case['mapped_roles']['status'] == 'MISSING_ATOMS'
    assert case['mapped_roles']['missing'] == {'N1': 'DOES_NOT_EXIST'}
    assert case['mapped_roles']['attachment_assignment'] == 'NOT_ASSESSED'


@pytest.mark.parametrize('damage,expected', [('identity', 'IDENTITY_OR_LAYOUT_REVIEW'),
                                            ('parse', 'IDENTITY_OR_LAYOUT_REVIEW'),
                                            ('extra_block', 'ASSESSMENT_REVIEW')])
def test_bad_sources_do_not_disappear_or_inherit_good_graph_checks(tmp_path, damage, expected):
    from nasolve.cif_catalogue import audit_catalogue
    book = workbook(tmp_path, [('thymine', row()), ('adenine', row('X99'))])
    ccd = tmp_path / 'ccd'; path = source(ccd)
    if damage == 'identity':
        path.write_text(path.read_text().replace('PDU', 'ZZZ'))
    elif damage == 'parse':
        path.write_text('data_bad\n_chem_comp.id "unfinished\n')
    else:
        path.write_text(path.read_text() + '\ndata_bad\n_chem_comp.id X\n_chem_comp.id X\n')
    report = audit_catalogue(book, ccd)
    cases = {c['component']: c for c in report['components']}
    assert len(cases) == 2
    assert cases['PDU']['source']['status'] == expected
    assert cases['PDU']['mapped_roles']['status'] == 'NOT_ASSESSED'
    assert cases['X99']['source']['status'] == 'MISSING'


def test_unsafe_id_and_symlink_escape_never_read_outside_anchor(tmp_path):
    from nasolve.cif_catalogue import audit_catalogue
    book = workbook(tmp_path, [('thymine', row('../PDU')), ('thymine', row())])
    ccd = tmp_path / 'ccd'; (ccd / 'p').mkdir(parents=True)
    outside = source(tmp_path / 'outside')
    (ccd / 'p' / 'data_PDU.cif').symlink_to(outside)
    original = Path.read_bytes
    def guarded(path):
        if path.resolve() == outside.resolve():
            raise AssertionError('escaped source anchor')
        return original(path)
    with patch.object(Path, 'read_bytes', guarded):
        report = audit_catalogue(book, ccd)
    assert [c['source']['status'] for c in report['components']] == ['INVALID_CODE', 'OUTSIDE_ANCHOR']


def test_relocation_original_gone_is_exact_and_read_only(tmp_path):
    from nasolve.cif_catalogue import audit_catalogue
    original = tmp_path / 'original'; original.mkdir()
    book = workbook(original, [('thymine', row())])
    source(original / 'ccd')
    report = audit_catalogue(book, original / 'ccd')
    moved = tmp_path / 'moved'; original.rename(moved)
    assert not original.exists()
    before = snapshot(moved)
    with patch('subprocess.Popen', side_effect=AssertionError('no processes')):
        assert audit_catalogue(moved / 'Ligands.xlsx', moved / 'ccd') == report
    assert snapshot(moved) == before


@pytest.mark.parametrize('target', ['source', 'workbook', 'missing'])
def test_mutation_during_audit_is_rejected(tmp_path, target):
    import nasolve.cif_catalogue as catalogue
    book = workbook(tmp_path, [('thymine', row()), ('adenine', row('X99'))])
    ccd = tmp_path / 'ccd'; path = source(ccd)
    original = catalogue._source
    def mutate(root, code):
        result = original(root, code)
        if code == 'X99':
            if target == 'missing':
                new = ccd / 'x' / 'data_X99.cif'; new.parent.mkdir(); new.write_text('changed')
            else:
                changed = path if target == 'source' else book
                changed.write_bytes(changed.read_bytes() + b'\n')
        return result
    with patch.object(catalogue, '_source', mutate), pytest.raises(ValueError, match='changed'):
        catalogue.audit_catalogue(book, ccd)


def test_loader_row_mismatch_is_rejected(tmp_path):
    from nasolve.cif_catalogue import audit_catalogue
    from restraints import residue_library
    book = workbook(tmp_path, [('thymine', row())]); ccd = tmp_path / 'ccd'; source(ccd)
    records = deepcopy(residue_library.load_residue_records(book)); records[0]['Ligand code'] = 'ZZZ'
    with patch.object(residue_library, 'load_residue_records', return_value=records), pytest.raises(ValueError, match='disagree'):
        audit_catalogue(book, ccd)


def test_cli_plain_text_and_missing_input_are_honest(tmp_path, capsys):
    book = workbook(tmp_path, [('thymine', row())]); ccd = tmp_path / 'ccd'; source(ccd)
    assert main(['cif-doctor', 'catalogue', '--workbook', str(book), '--ccd-root', str(ccd)]) == 2
    text = capsys.readouterr().out
    assert '1 records; 1 distinct codes' in text
    assert 'Native qualification: not assessed' in text
    assert main(['cif-doctor', 'catalogue', '--workbook', str(tmp_path / 'absent.xlsx'),
                 '--ccd-root', str(ccd), '--json']) == 1
    assert 'CIF Doctor:' in capsys.readouterr().err


def test_absent_role_marker_and_empty_role_record_do_not_claim_missing_atoms(tmp_path):
    from nasolve.cif_catalogue import audit_catalogue
    book = workbook(tmp_path, [('thymine', row(C5='/')), ('unique', row('ATL', N1=None, C5=None))])
    ccd = tmp_path / 'ccd'; source(ccd); source(ccd, 'ATL')
    cases = {c['component']: c for c in audit_catalogue(book, ccd)['components']}
    assert cases['PDU']['mapped_roles']['status'] == 'ATOM_NAMES_PRESENT'
    assert cases['PDU']['mapped_roles']['absent_role_markers'] == {'C5': '/'}
    assert cases['PDU']['workbook_rows'][0]['operational_signature']['atom_roles']['C5'] == '/'
    assert cases['ATL']['mapped_roles']['status'] == 'NO_MAPPED_ROLES'


def test_spelling_conflict_remains_explicit_and_label_groups_overlap(tmp_path):
    from nasolve.cif_catalogue import audit_catalogue
    book = workbook(tmp_path, [('thymine', row(**{'Sugar Type': "2'mod"})),
                               ('thymine', row(**{'Sugar Type': "2'-mod"}))])
    ccd = tmp_path / 'ccd'; source(ccd)
    report = audit_catalogue(book, ccd)
    assert report['components'][0]['record_conflicts'] == [
        {'section': 'context', 'field': 'Sugar Type', 'values_in_workbook_order': ["2'mod", "2'-mod"]}]
    assert len(report['sugar_label_coverage']) == 2
    assert all(g['records'] == g['distinct_codes'] == 1 for g in report['sugar_label_coverage'])


@pytest.mark.parametrize('damage', ['duplicate_header', 'missing_sheet', 'no_records', 'corrupt_zip'])
def test_bad_workbook_is_a_clear_input_error(tmp_path, damage, capsys):
    book = workbook(tmp_path, [('thymine', row())])
    if damage == 'corrupt_zip':
        book.write_bytes(b'not a workbook')
    else:
        wb = openpyxl.load_workbook(book)
        if damage == 'duplicate_header': wb['thymine']['B1'] = 'Ligand code'
        if damage == 'missing_sheet': del wb['adenine']
        if damage == 'no_records': wb['thymine'].delete_rows(2)
        wb.save(book); wb.close()
    ccd = tmp_path / 'ccd'; source(ccd)
    assert main(['cif-doctor', 'catalogue', '--workbook', str(book), '--ccd-root', str(ccd)]) == 1
    assert 'CIF Doctor:' in capsys.readouterr().err


def test_five_character_exact_code_is_not_truncated(tmp_path):
    from nasolve.cif_catalogue import audit_catalogue
    book = workbook(tmp_path, [('thymine', row('A1AAZ'))])
    ccd = tmp_path / 'ccd'; original = source(ccd)
    target = ccd / 'a' / 'data_A1AAZ.cif'; target.parent.mkdir()
    target.write_text(original.read_text().replace('PDU', 'A1AAZ'))
    case = audit_catalogue(book, ccd)['components'][0]
    assert case['source']['status'] == 'RECORDED'
    assert case['source']['reference']['path'] == 'a/data_A1AAZ.cif'


def test_broken_symlink_stays_visible_without_false_mutation(tmp_path):
    from nasolve.cif_catalogue import audit_catalogue
    book = workbook(tmp_path, [('thymine', row())]); ccd = tmp_path / 'ccd'
    (ccd / 'p').mkdir(parents=True); (ccd / 'p' / 'data_PDU.cif').symlink_to(ccd / 'missing.cif')
    assert audit_catalogue(book, ccd)['components'][0]['source']['status'] == 'INVALID_RESOURCE'
