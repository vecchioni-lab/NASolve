"""Read-only coverage census of an explicit NARestraints workbook and CCD root.

This is source and role evidence, not native qualification or a repair policy.
Workbook labels never establish linkage chemistry or authorize atom deletion.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from hashlib import sha256
from pathlib import Path
import json
import re
from zipfile import BadZipFile

import openpyxl
from restraints import residue_library

from .cif_doctor import assess_dictionaries
from .cif_linked_preparation import component_values
from .component_records import resolve_component_record, role_signature

POLICY_ID = "cif-narestraints-catalogue-census-v1"
MAX_SOURCE_BYTES = 8 * 1024 * 1024


def _digest(data):
    return {"sha256": sha256(data).hexdigest(), "size": len(data)}


def _workbook_records(path):
    """Use NAR's actual loader; independently bind its output to worksheet rows."""
    locations = []
    try:
        workbook = openpyxl.load_workbook(path, read_only=True)
    except (BadZipFile, openpyxl.utils.exceptions.InvalidFileException) as exc:
        raise ValueError("Require a readable NARestraints XLSX workbook") from exc
    try:
        for name in residue_library.AUTHORITATIVE_SHEETS:
            sheet = workbook[name]
            populated = ((number, row) for number, row in enumerate(sheet.values, 1)
                         if any(cell is not None for cell in row))
            _, headers = next(populated)
            named = [h for h in headers if h is not None]
            if len(named) != len(set(named)) or "Ligand code" not in named or "Function" not in named:
                raise ValueError("Ambiguous NARestraints headers in " + name)
            code_index = headers.index("Ligand code")
            for number, row in populated:
                if row[0] is not None:
                    locations.append({"sheet": name, "row": number, "code": str(row[code_index])})
    except (KeyError, StopIteration, IndexError) as exc:
        raise ValueError("Incomplete NARestraints workbook layout") from exc
    finally:
        workbook.close()
    try:
        records = residue_library.load_residue_records(path)
    except (KeyError, IndexError) as exc:
        raise ValueError("Incomplete NARestraints workbook records") from exc
    if not records:
        raise ValueError("NARestraints workbook contains no ligand records")
    if len(locations) != len(records) or any(
        loc["code"] != rec["Ligand code"] or loc["sheet"] != rec["Source sheet"]
        for loc, rec in zip(locations, records)
    ):
        raise ValueError("NARestraints loader and worksheet row inventory disagree")
    # Fail explicitly on non-JSON cell types instead of losing their provenance.
    try:
        json.dumps(records, allow_nan=False)
    except (TypeError, ValueError) as exc:
        raise ValueError("Unsupported NARestraints workbook cell value") from exc
    return records, locations


def _source(root, code):
    if not re.fullmatch(r"[A-Z0-9]{1,5}", code):
        return {"status": "INVALID_CODE", "graph_status": "NOT_ASSESSED"}, None
    relative = f"{code[0].lower()}/data_{code}.cif"
    ref = {"anchor": "ccd_directory", "path": relative}
    path = root / relative
    result = {"reference": ref, "status": "MISSING", "graph_status": "NOT_ASSESSED"}
    # Reject a symlink escape before reading any resource outside the chosen root.
    if not path.resolve().is_relative_to(root):
        result["status"] = "OUTSIDE_ANCHOR"
        return result, None
    if not path.exists():
        if path.is_symlink():
            result["status"] = "INVALID_RESOURCE"
        return result, None
    if not path.is_file() or path.stat().st_size > MAX_SOURCE_BYTES:
        result["status"] = "INVALID_RESOURCE"
        return result, None
    try:
        data = path.read_bytes()
    except OSError as exc:
        result.update(status="UNREADABLE", error_type=type(exc).__name__)
        return result, None
    ref.update(_digest(data))
    assessment = assess_dictionaries([path], stage="source")
    if assessment["input_integrity"] != "UNCHANGED" or assessment["inputs"][0]["sha256"] != ref["sha256"]:
        raise ValueError("CCD source changed during catalogue assessment: " + code)
    result.update(status="RECORDED", findings=assessment["findings"])
    components = assessment["components"]
    if len(components) != 1 or components[0]["component_id"] != code:
        result.update(status="IDENTITY_OR_LAYOUT_REVIEW", graph_status="NOT_ASSESSED")
        return result, None
    component = components[0]
    result.update(graph_status=component["verdicts"]["graph"], verdicts=component["verdicts"])
    # File-level parse failures must not be hidden by a valid component block.
    if any(f["status"] in {"FAIL", "ERROR"} for f in assessment["findings"]):
        result["status"] = "ASSESSMENT_REVIEW"
    if result["status"] != "RECORDED" or result["graph_status"] != "PASS":
        return result, None
    try:
        values = component_values(path, code)
    except ValueError as exc:
        result.update(status="METADATA_REVIEW", metadata_error=str(exc))
        return result, None
    inventory = component["inventory"]
    atoms = {a["atom_id"]: a for a in inventory["atoms"]}
    neighbors = {a: set() for a in atoms}
    for bond in inventory["bond_orders"]:
        first, second = bond["atoms"]
        neighbors[first].add(second)
        neighbors[second].add(first)
    result["declarations"] = {
        "type": values.get("_chem_comp.type", values.get("_chem_comp.group", [None]))[0],
        "formal_charge": values.get("_chem_comp.pdbx_formal_charge", [None])[0],
        "atom_count": inventory["atom_count"], "bond_count": inventory["bond_count"],
        "phosphorus_atoms": [{"atom_id": name, "neighbors": sorted(neighbors[name])}
                             for name, atom in sorted(atoms.items()) if atom["element"] == "P"],
        "op3": None if "OP3" not in atoms else {
            "element": atoms["OP3"]["element"], "neighbors": sorted(neighbors["OP3"])},
        "stereo_markers": {name: atom["stereo_declarations"] for name, atom in sorted(atoms.items())
                           if atom["stereo_declarations"]},
    }
    return result, (atoms, neighbors)


def audit_catalogue(workbook, ccd_root):
    """Census every loader record and exact code; write nothing and call no SDK.

    Both anchors must be supplied explicitly on every invocation. The output is
    reproducible after relocation and binds source bytes, not an installed path.
    Missing resources, ambiguous records and unavailable role checks stay visible.
    No native support or absence of chemical defects is inferred from this audit.
    """
    workbook = Path(workbook).expanduser().resolve(strict=True)
    root = Path(ccd_root).expanduser().resolve(strict=True)
    if not workbook.is_file() or not root.is_dir():
        raise ValueError("Require an explicit workbook file and CCD directory")
    before = workbook.read_bytes()
    records, locations = _workbook_records(workbook)
    grouped = defaultdict(list)
    for record, location in zip(records, locations):
        grouped[record["Ligand code"]].append((record, location))
    components = []
    for code, entries in sorted(grouped.items()):
        matching = [rec for rec, _ in entries]
        row = {"component": code, "workbook_rows": [
            {**loc, "labels": {k: rec.get(k) for k in ("Base Analog", "Sugar Type", "Phosphate")},
             "operational_signature": role_signature(rec)} for rec, loc in entries],
            "native_qualification": "NOT_ASSESSED", "linkage_profile": "NOT_ASSIGNED",
            "automatic_adoption": False}
        try:
            selected, resolution = resolve_component_record(matching, code)
            row["record_resolution"] = resolution["resolution"]
        except ValueError:
            selected = None
            row["record_resolution"] = "CONFLICTING_OPERATIONAL_RECORDS"
            signatures = [role_signature(rec) for rec in matching]
            row["record_conflicts"] = [
                {"section": section, "field": field,
                 "values_in_workbook_order": [sig[section].get(field) for sig in signatures]}
                for section in ("context", "atom_roles")
                for field in sorted(set().union(*(sig[section] for sig in signatures)))
                if len({json.dumps(sig[section].get(field), sort_keys=True) for sig in signatures}) > 1
            ]
        source, graph = _source(root, code)
        row["source"] = source
        row["mapped_roles"] = {"status": "NOT_ASSESSED"}
        if selected is not None and graph is not None:
            atoms, neighbors = graph
            raw_roles = role_signature(selected)["atom_roles"]
            # NAR validator/mirroring explicitly use '/' as an absent-role marker.
            # Retain it in the raw signature; do not invent an atom or edit NAR.
            placeholders = {role: atom for role, atom in raw_roles.items() if atom == "/"}
            roles = {role: atom for role, atom in raw_roles.items() if atom != "/"}
            missing = {role: atom for role, atom in roles.items() if atom not in atoms}
            base = set(roles.values()) & atoms.keys()
            boundaries = sorted({tuple(sorted((a, b))) for a in base for b in neighbors[a]
                                 if b not in base and atoms[b]["element"] not in {"H", "D"}})
            row["mapped_roles"] = {
                "status": "NO_MAPPED_ROLES" if not roles else "MISSING_ATOMS" if missing else "ATOM_NAMES_PRESENT",
                "missing": missing, "absent_role_markers": placeholders, "mapped_atom_count": len(base),
                "base_to_nonmapped_heavy_bonds": [list(pair) for pair in boundaries],
                "glycosidic_atom_from_loader": selected.get("Glycosidic atom"),
                "attachment_assignment": "NOT_ASSESSED",
            }
        actions = []
        if selected is None:
            actions.append("RECONCILE_OPERATIONAL_RECORDS")
        if source["status"] == "MISSING":
            actions.append("ACQUIRE_EXACT_SOURCE")
        elif graph is None:
            actions.append("REVIEW_SOURCE")
        if row["mapped_roles"]["status"] in {"MISSING_ATOMS", "NO_MAPPED_ROLES"}:
            actions.append("RECONCILE_MAPPED_ATOMS")
        row["next_actions"] = actions + ["ESTABLISH_ATTACHMENT_PROFILE", "QUALIFY_NATIVE_CONSUMER"]
        components.append(row)
    if workbook.read_bytes() != before:
        raise ValueError("NARestraints workbook changed during catalogue assessment")
    # Check the whole source set again, including initially absent resources.
    for component in components:
        source = component["source"]
        ref = source.get("reference")
        if not ref:
            continue
        path = root / ref["path"]
        if "sha256" in ref:
            if not path.resolve().is_relative_to(root) or _digest(path.read_bytes()) != {k: ref[k] for k in ("sha256", "size")}:
                raise ValueError("CCD source changed during catalogue assessment: " + component["component"])
        elif source["status"] == "MISSING" and (path.exists() or path.is_symlink()):
            raise ValueError("Missing CCD source changed during catalogue assessment: " + component["component"])
    def count(values):
        return dict(sorted(Counter(values).items()))
    def labels(field):
        counts = Counter(rec.get(field) for rec in records)
        return [{"value": value, "records": number} for value, number in sorted(
            counts.items(), key=lambda item: (item[0] is not None, str(item[0])))]
    return {
        "schema_version": 1, "policy_id": POLICY_ID, "status": "REVIEW",
        "audit_kind": "narestraints_catalogue_census", "coverage": "COMPLETE_FOR_SELECTED_WORKBOOK",
        "workbook": {"anchor": "workbook_directory", "path": workbook.name, **_digest(before)},
        "source_layout": "ccd_directory/{first_character_lower}/data_{exact_code}.cif",
        "loader": {"module": "restraints.residue_library", "openpyxl_version": openpyxl.__version__,
                   **_digest(Path(residue_library.__file__).read_bytes())},
        "authoritative_sheets": list(residue_library.AUTHORITATIVE_SHEETS),
        "summary": {
            "records": len(records), "distinct_codes": len(components),
            "records_by_sheet": count(loc["sheet"] for loc in locations),
            "sugar_labels": labels("Sugar Type"), "phosphate_labels": labels("Phosphate"),
            "record_resolution": count(c["record_resolution"] for c in components),
            "source_status": count(c["source"]["status"] for c in components),
            "graph_status": count(c["source"]["graph_status"] for c in components),
            "mapped_roles": count(c["mapped_roles"]["status"] for c in components),
            "native_qualification": {"NOT_ASSESSED": len(components)},
        },
        "sugar_label_coverage": [
            {"value": label, "records": sum(rec.get("Sugar Type") == label for rec in records),
             "distinct_codes": len(cases),
             "source_status": count(c["source"]["status"] for c in cases),
             "record_resolution": count(c["record_resolution"] for c in cases),
             "mapped_roles": count(c["mapped_roles"]["status"] for c in cases)}
            for label in sorted({rec.get("Sugar Type") for rec in records}, key=lambda value: (value is not None, str(value)))
            for cases in [[c for c in components if any(r["labels"]["Sugar Type"] == label for r in c["workbook_rows"])]]
        ],
        "components": components, "input_integrity": "UNCHANGED", "native_execution": "NOT_RUN",
        "transformations": [], "automatic_adoption": False,
        "limitations": [
            "Workbook membership and labels are not graph, linkage, stereochemical or numerical authority.",
            "Blank phosphate metadata does not establish phosphate absence; OP3 removal needs explicit site context.",
            "Mapped atom presence does not validate base roles, glycosidic assignment or pair/stack compilation.",
            "Exact sugar-label groups may overlap for conflicting duplicate codes; spelling is not normalized.",
            "Boundary bonds include substituents; they do not assign attachment sites or leaving groups.",
            "This census does not compile NAR recipes, generate restraints, qualify native consumers or select models.",
            "NOT_ASSESSED counts describe this audit, not the absence of prior ligand-specific qualification.",
        ],
    }
